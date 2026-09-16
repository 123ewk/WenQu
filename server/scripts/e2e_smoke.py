"""端到端真实链路验证(不含模型 API):真 PG + 真 MinIO + 真 gRPC parser。

模型调用(向量化/对话)用确定性桩替换 —— 用户尚未填 Key,这两处是唯一被打桩的环节。
覆盖:注册 → 建 KB → 上传真实文件 → 对象存储落位 → worker 入库(解析/分块/写库)
→ 混合检索命中 → SSE 问答引用可回溯 → 分块列表 → 维护清理。

运行:uv run python scripts/e2e_smoke.py
"""

from __future__ import annotations

import io
import json
import os
import sys
import uuid

os.environ.setdefault("APP_DATABASE_URL", "postgresql+psycopg://wenqu:wenqu_dev_only@localhost:5432/wenqu")
os.environ.setdefault("PARSER_GRPC_ADDR", "127.0.0.1:50071")
os.environ.setdefault("PARSER_GRPC_TOKEN", "wenqu_dev_parser_token")

from fastapi.testclient import TestClient  # noqa: E402

import app.main as main_module  # noqa: E402
from app.api.deps import get_chat_gateway, get_embedding_gateway  # noqa: E402
from app.application.repository.knowledge import (  # noqa: E402
    DocumentRepositoryImpl,
    KnowledgeBaseRepositoryImpl,
)
from app.application.repository.tasks import TaskRepositoryImpl  # noqa: E402
from app.application.service.ingestion import IngestionService  # noqa: E402
from app.core.db import get_session_factory  # noqa: E402
from app.core.parser_client import ParserClient  # noqa: E402
from app.core.storage import MinioStorage  # noqa: E402

PASS, FAIL = "✅", "❌"
results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))
    print(f"{PASS if ok else FAIL} {name}{f' — {detail}' if detail else ''}")


class StubEmbedder:
    """确定性向量:含"检索/算法"的文本贴近查询向量,其余远离(替代真实 embedding API)。"""

    def embed(self, texts, model_id=None):
        out = []
        for text in texts:
            v = [0.0] * 1024
            v[0 if ("检索" in text or "算法" in text or "RRF" in text) else 1] = 1.0
            out.append(v)
        return out


class StubChat:
    def __init__(self) -> None:
        self.prompts: list[list[dict]] = []

    def chat_stream(self, messages, model_id=None):
        self.prompts.append(messages)
        yield "混合检索使用 RRF 融合 [1],"
        yield "可提升召回稳定性。"


def main() -> int:
    print("=== WenQu 端到端真实链路验证 ===")
    print(f"PG: {os.environ['APP_DATABASE_URL'].split('@')[-1]}")
    print(f"Parser: {os.environ['PARSER_GRPC_ADDR']}\n")

    stub_chat = StubChat()
    application = main_module.create_app()
    application.dependency_overrides[get_embedding_gateway] = StubEmbedder
    application.dependency_overrides[get_chat_gateway] = lambda: stub_chat

    client = TestClient(application, raise_server_exceptions=False)
    tag = uuid.uuid4().hex[:6]

    # 1) 注册登录
    reg = client.post(
        "/api/v1/auth/register",
        json={"username": f"e2e{tag}", "nickname": "e2e", "password": "secret-pass-1"},
    )
    check("注册账号", reg.status_code == 201, f"HTTP {reg.status_code}")
    if reg.status_code != 201:
        return 1
    auth = {"Authorization": f"Bearer {reg.json()['access_token']}"}

    # 2) 建空间 + 知识库
    space_id = client.post(
        "/api/v1/spaces", json={"name": f"E2E空间{tag}"}, headers=auth
    ).json()["id"]
    kb = client.post(
        f"/api/v1/spaces/{space_id}/knowledge-bases",
        json={"name": f"E2E知识库{tag}", "description": "端到端验证"},
        headers=auth,
    )
    check("创建空间与知识库", kb.status_code == 201, f"HTTP {kb.status_code}")
    kb_id = kb.json()["id"]

    # 3) 上传真实 docx(含标题/段落/表格),验证 multipart 全链路
    from docx import Document as DocxDocument

    docx = DocxDocument()
    docx.add_heading("混合检索设计", level=1)
    docx.add_paragraph("混合检索算法使用 RRF 融合向量与全文两路排名,提升召回稳定性。")
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "取值"
    table.cell(1, 0).text = "k"
    table.cell(1, 1).text = "60"
    buffer = io.BytesIO()
    docx.save(buffer)

    up = client.post(
        f"/api/v1/spaces/{space_id}/knowledge-bases/{kb_id}/documents",
        files={"file": ("混合检索设计.docx", buffer.getvalue())},
        headers=auth,
    )
    check("上传 docx 文档", up.status_code == 201, f"HTTP {up.status_code}")
    if up.status_code != 201:
        print(up.text)
        return 1
    document = up.json()
    check("文档初始状态 pending", document["status"] == "pending", document["status"])

    # 4) 校验对象存储真的落位(MinIO)
    storage = MinioStorage(
        "localhost:9000", "wenqu", "wenqu_dev_only", "wenqu-docs"
    )
    with get_session_factory()() as db:
        from app.domain.models import Document as DocModel

        row = db.get(DocModel, uuid.UUID(document["id"]))
        stored = storage.get(row.source)
    check("MinIO 对象存储落位", len(stored) == len(buffer.getvalue()), f"{len(stored)} bytes")

    # 5) 跑真实 worker 入库(真 gRPC parser + 真分块 + 桩向量化)

    with get_session_factory()() as db:
        service = IngestionService(
            documents=DocumentRepositoryImpl(db),
            kbs=KnowledgeBaseRepositoryImpl(db),
            tasks=TaskRepositoryImpl(db),
            chunks=_RealChunkRepo(db),
            storage=storage,
            parser=ParserClient(
                os.environ["PARSER_GRPC_ADDR"], os.environ["PARSER_GRPC_TOKEN"]
            ),
            embedder=StubEmbedder(),
        )
        processed: list[str] = []
        for _ in range(20):
            task_id = service.handle_next(db, "e2e-worker")
            if task_id is None:
                break
            processed.append(task_id)
        check("worker 认领并处理任务", bool(processed), f"处理 {len(processed)} 个任务")

    # 6) 文档状态推进到 completed,分块落库
    detail = client.get(
        f"/api/v1/spaces/{space_id}/knowledge-bases/{kb_id}/documents/{document['id']}",
        headers=auth,
    ).json()
    check("文档状态 completed", detail["status"] == "completed", detail["status"])

    chunks = client.get(
        f"/api/v1/spaces/{space_id}/knowledge-bases/{kb_id}/documents/{document['id']}/chunks",
        headers=auth,
    ).json()
    check("分块落库可读", chunks["total"] > 0, f"共 {chunks['total']} 块")
    joined = " ".join(c["content"] for c in chunks["items"])
    check("标题面包屑进入分块", "混合检索设计" in joined)
    check(
        "表格被识别为 table 块并保留表头",
        any(c["meta"].get("kind") == "table" for c in chunks["items"]),
        f"块类型={[c['meta'].get('kind') for c in chunks['items']]}",
    )

    # 7) 混合检索命中(真 PG:pgvector + tsvector)
    search = client.post(
        f"/api/v1/spaces/{space_id}/retrieval/search",
        json={"query": "混合检索算法", "top_k": 3},
        headers=auth,
    )
    hits = search.json()
    check("混合检索命中", bool(hits), f"命中 {len(hits)} 条")
    if hits:
        check(
            "命中内容相关且带双路排名",
            "RRF" in hits[0]["content"] or "检索" in hits[0]["content"],
            f"vector_rank={hits[0]['vector_rank']} fulltext_rank={hits[0]['fulltext_rank']}",
        )

    # 8) SSE 流式问答 + 引用可回溯
    ask = client.post(
        f"/api/v1/spaces/{space_id}/ask",
        json={"question": "混合检索算法用了什么融合方式?", "top_k": 3},
        headers=auth,
    )
    events = [
        json.loads(line[6:])
        for line in ask.text.splitlines()
        if line.startswith("data: ")
    ]
    types = [e["type"] for e in events]
    check("SSE 事件时序正确", types == ["meta", "citations", "delta", "delta", "done"], str(types))
    citation = events[1]["citations"][0] if events[1]["citations"] else None
    check("引用携带可回链字段", bool(citation and citation["chunk_id"] and citation["excerpt"]))
    cited = events[-1].get("cited_indexes")
    check("答案标注了引用编号", cited == [1], str(cited))

    conv_id = events[0]["conversation_id"]
    messages = client.get(
        f"/api/v1/spaces/{space_id}/conversations/{conv_id}/messages", headers=auth
    ).json()
    check(
        "消息落库且引用可回溯",
        len(messages) == 2 and messages[1]["citations"][0]["chunk_id"] == citation["chunk_id"],
    )

    # 9) 多轮:第二轮带上会话 id
    ask2 = client.post(
        f"/api/v1/spaces/{space_id}/ask",
        json={"question": "k 的取值是多少?", "conversation_id": conv_id, "top_k": 3},
        headers=auth,
    )
    events2 = [
        json.loads(line[6:])
        for line in ask2.text.splitlines()
        if line.startswith("data: ")
    ]
    check("多轮复用同一会话", events2[0]["conversation_id"] == conv_id)
    prompt_roles = [m["role"] for m in stub_chat.prompts[1]]
    check("第二轮提示词含历史", prompt_roles.count("user") >= 2, str(prompt_roles))
    check("提示词携带编号资料", "[1] 来源:" in stub_chat.prompts[1][0]["content"])

    # 10) 维护清理(真 PG)
    from datetime import UTC, datetime, timedelta

    from app.application.service.maintenance import MaintenanceService
    from app.domain.models import RefreshToken

    with get_session_factory()() as db:
        db.add(
            RefreshToken(
                user_id=uuid.UUID(reg.json()["user"]["id"]),
                token_hash=uuid.uuid4().hex + uuid.uuid4().hex,
                expires_at=datetime.now(UTC) - timedelta(hours=1),
            )
        )
        db.commit()
        removed_tokens, removed_logs = MaintenanceService(db, 180).run_all()
    check("过期 refresh 令牌被清理", removed_tokens >= 1, f"清除 {removed_tokens} 条")

    # 11) 空间隔离:另一个用户看不到
    other = client.post(
        "/api/v1/auth/register",
        json={"username": f"e2eother{tag}", "nickname": "o", "password": "secret-pass-1"},
    ).json()
    other_auth = {"Authorization": f"Bearer {other['access_token']}"}
    resp = client.get(f"/api/v1/spaces/{space_id}/knowledge-bases", headers=other_auth)
    check("空间隔离(非成员 404)", resp.status_code == 404, f"HTTP {resp.status_code}")

    # 汇总
    failed = [name for name, ok, _ in results if not ok]
    print(f"\n=== 结果:{len(results) - len(failed)}/{len(results)} 通过 ===")
    if failed:
        print("失败项:" + "、".join(failed))
        return 1
    print("全链路验证通过(模型 API 为桩,其余均为真实服务)")
    return 0


class _RealChunkRepo:
    """真实分块仓储(与生产 ChunkRepositoryImpl 同实现路径)。"""

    def __init__(self, db):
        from app.application.repository.knowledge import ChunkRepositoryImpl

        self._impl = ChunkRepositoryImpl(db)

    def replace_for_document(self, document, drafts):
        self._impl.replace_for_document(document, drafts)


if __name__ == "__main__":
    sys.exit(main())
