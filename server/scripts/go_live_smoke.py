"""Go-Live 冒烟七条(prod 形态真实 HTTP 链路,经 web/nginx 入口)。

与 e2e_smoke.py 的区别:这里不打进程内 TestClient、不打桩模型 ——
全部走 `WENQU_BASE_URL`(默认 http://web:8080,即生产唯一入口),
真实 worker 容器异步入库、真实模型 Key(embedding/chat),覆盖:
条1 四格式上传→completed / 条2 损坏文件失败+reparse / 条3 引用+删库兜底 /
条4 断流按 seq 续播不重不丢 / 条5 受限 API Key 范围外 403 / 条6 Viewer 403+非成员 404+审计。
条7(restart server 任务重认领)由外部执行:docker compose restart server 后查
tasks 表无 processing 残留 + /health 通过。

运行(演练栈容器内):docker compose exec server python scripts/go_live_smoke.py
环境变量:WENQU_BASE_URL;真实模型 Key(DASHSCOPE_API_KEY/DEEPSEEK_API_KEY)。
OPT-25 起内置模型下线:冒烟要求库内已配 chat+embedding 默认模型;
未配置时可用 WENQU_ADMIN_TOKEN(superuser JWT)让本脚本经管理端 API 自动配置。
"""

from __future__ import annotations

import io
import json
import os
import time
import uuid

import httpx

BASE = os.environ.get("WENQU_BASE_URL", "http://web:8080").rstrip("/")
TIMEOUT = httpx.Timeout(60.0, connect=10.0)
PASS, FAIL = "✅", "❌"
results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))
    print(f"{PASS if ok else FAIL} {name}{f' — {detail}' if detail else ''}", flush=True)


def ok_req(client: httpx.Client, method: str, url: str, tries: int = 4, **kw) -> httpx.Response:
    # 演练实测:建档/建库后秒级发起的下一个请求偶发 404(手动 curl 复现不了)。
    # 对"预期成功"的调用退避重试,把可见性竞态与稳定缺陷区分开;重试只打印不判分。
    r: httpx.Response | None = None
    for i in range(tries):
        r = client.request(method, url, **kw)
        if r.status_code // 100 == 2:
            return r
        print(f"⚠️ {method} {url} 第{i + 1}次 HTTP {r.status_code} {r.text[:120]}", flush=True)
        time.sleep(1 + i)
    assert r is not None
    return r


def sse_events(resp: httpx.Response) -> list[dict]:
    return [json.loads(line[6:]) for line in resp.text.splitlines() if line.startswith("data: ")]


def poll_doc(
    client: httpx.Client, auth: dict, space: str, kb: str, doc: str, timeout_s: int = 180
) -> dict:
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        d = client.get(
            f"/api/v1/spaces/{space}/knowledge-bases/{kb}/documents/{doc}", headers=auth
        ).json()
        if d["status"] in ("completed", "failed"):
            return d
        time.sleep(2)
    raise AssertionError(f"document {doc} 超时未达终态")


def make_files() -> dict[str, tuple[str, bytes]]:
    tag = uuid.uuid4().hex[:6]
    import openpyxl
    import pymupdf
    from docx import Document as DocxDocument

    docx_b = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading(f"混合检索设计 {tag}", level=1)
    doc.add_paragraph("混合检索算法使用 RRF 融合向量与全文两路排名,提升召回稳定性。参数 k 取 60。")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "参数"
    t.cell(0, 1).text = "取值"
    t.cell(1, 0).text = "k"
    t.cell(1, 1).text = "60"
    doc.save(docx_b)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["主题", "说明"])
    ws.append(["RRF", "向量与全文双路排名融合"])
    xlsx_b = io.BytesIO()
    wb.save(xlsx_b)

    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page()
    page.insert_text((72, 72), "Hybrid retrieval fuses vector and fulltext ranks by RRF.")
    pdf_b = io.BytesIO(pdf_doc.tobytes(garbage=4))
    pdf_doc.close()

    md = (
        f"# 向量索引 {tag}\n\n向量检索把文本编码为 embedding,用近邻召回。RRF 融合提升稳定性。\n"
    ).encode()
    return {
        "docx": (f"混合检索设计_{tag}.docx", docx_b.getvalue()),
        "xlsx": (f"表格说明_{tag}.xlsx", xlsx_b.getvalue()),
        "pdf": (f"retrieval_{tag}.pdf", pdf_b.getvalue()),
        "md": (f"向量索引_{tag}.md", md),
    }


def ensure_models(client: httpx.Client, auth: dict, tag: str) -> tuple[bool, str]:
    """OPT-25:内置模型下线,冒烟链路要先有 chat+embedding 默认模型。

    已配置 → 直接过;未配置且有 WENQU_ADMIN_TOKEN → 经管理端 API 现配
    (供应商/模型 id 带 tag,重复运行互不冲突);都没有 → 判失败。
    """
    catalog = client.get("/api/v1/models", headers=auth).json()
    if catalog["defaults"].get("chat") and catalog["defaults"].get("embedding"):
        return True, f"已配置:{catalog['defaults']}"
    token = os.environ.get("WENQU_ADMIN_TOKEN", "")
    dash_key = os.environ.get("DASHSCOPE_API_KEY", "")
    ds_key = os.environ.get("DEEPSEEK_API_KEY", "")
    if not token or not dash_key or not ds_key:
        return False, (
            "库内无默认模型,且缺 WENQU_ADMIN_TOKEN/DASHSCOPE_API_KEY/DEEPSEEK_API_KEY,无法自动配置"
        )
    ahdr = {"Authorization": f"Bearer {token}"}
    plans = [
        (
            {
                "name": f"smoke-dashscope-{tag}",
                "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
                "api_key": dash_key,
            },
            {
                "model_id": f"smoke-dashscope/text-embedding-v3-{tag}",
                "kind": "embedding",
                "model_name": "text-embedding-v3",
                "dims": 1024,
                "is_default": True,
            },
        ),
        (
            {
                "name": f"smoke-deepseek-{tag}",
                "base_url": "https://api.deepseek.com/v1",
                "api_key": ds_key,
            },
            {
                "model_id": f"smoke-deepseek/deepseek-chat-{tag}",
                "kind": "chat",
                "model_name": "deepseek-chat",
                "is_default": True,
            },
        ),
    ]
    for provider_in, model_in in plans:
        pr = client.post("/api/v1/admin/model-providers", json=provider_in, headers=ahdr)
        if pr.status_code != 201:
            return False, f"建供应商失败 HTTP {pr.status_code} {pr.text[:120]}"
        mr = client.post(
            "/api/v1/admin/models",
            json={**model_in, "provider_id": pr.json()["id"]},
            headers=ahdr,
        )
        if mr.status_code != 201:
            return False, f"建模型失败 HTTP {mr.status_code} {mr.text[:120]}"
    return True, f"经管理端 API 配置完成:{tag}"


def main() -> int:
    print(f"=== WenQu 冒烟七条(prod 形态,HTTP 真实链路)@ {BASE} ===")
    tag = uuid.uuid4().hex[:6]
    client = httpx.Client(base_url=BASE, timeout=TIMEOUT)

    r = ok_req(
        client,
        "POST",
        "/api/v1/auth/register",
        json={"username": f"smoke{tag}", "nickname": "冒烟", "password": "secret-pass-1"},
    )
    check("注册用户", r.status_code == 201, f"HTTP {r.status_code}")
    auth = {"Authorization": f"Bearer {r.json()['access_token']}"}
    me_id = r.json()["user"]["id"]
    mk_ok, mk_detail = ensure_models(client, auth, tag)
    check("前置:库内 chat+embedding 默认模型可用(OPT-25)", mk_ok, mk_detail)
    if not mk_ok:
        print("模型未就绪,后续条目必挂,提前退出。")
        return 1
    space = ok_req(
        client, "POST", "/api/v1/spaces", json={"name": f"冒烟空间{tag}"}, headers=auth
    ).json()["id"]
    kb_a = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/knowledge-bases",
        json={"name": f"KB-A{tag}", "description": "四格式"},
        headers=auth,
    ).json()["id"]
    kb_b = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/knowledge-bases",
        json={"name": f"KB-B{tag}", "description": "空库(兜底用)"},
        headers=auth,
    ).json()["id"]

    # ---- 条 1:PDF/Office/Markdown/表格 各传一篇 → 真 worker 异步入库 ----
    statuses: dict[str, str] = {}
    doc_ids: dict[str, str] = {}
    for fmt, (fname, data) in make_files().items():
        up = ok_req(
            client,
            "POST",
            f"/api/v1/spaces/{space}/knowledge-bases/{kb_a}/documents",
            files={"file": (fname, data)},
            headers=auth,
        )
        doc_ids[fmt] = up.json()["id"]
        statuses[fmt] = up.json()["status"]
    check(
        "条1 四格式上传受理(pending)", all(s == "pending" for s in statuses.values()), str(statuses)
    )
    finals = {fmt: poll_doc(client, auth, space, kb_a, did) for fmt, did in doc_ids.items()}
    bad = {f: d["status"] for f, d in finals.items() if d["status"] != "completed"}
    check("条1 真 worker 全部 completed(含向量化)", not bad, str(bad) if bad else "4/4 completed")
    chunk_totals = {
        f: client.get(
            f"/api/v1/spaces/{space}/knowledge-bases/{kb_a}/documents/{d['id']}/chunks",
            headers=auth,
        ).json()["total"]
        for f, d in finals.items()
    }
    check("条1 四格式分块数均 >0", all(v > 0 for v in chunk_totals.values()), str(chunk_totals))

    # ---- 条 2:损坏 docx → failed + error_message 可见 + reparse 可重试 ----
    broken = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/knowledge-bases/{kb_a}/documents",
        files={"file": ("broken.docx", b"PK\x03\x04not-a-real-zip" + os.urandom(64))},
        headers=auth,
    ).json()
    bd = poll_doc(client, auth, space, kb_a, broken["id"], timeout_s=120)
    check("条2 损坏文件状态 failed", bd["status"] == "failed", bd["status"])
    check(
        "条2 error_message 可见", bool(bd.get("error_message")), str(bd.get("error_message"))[:80]
    )
    rp = client.post(
        f"/api/v1/spaces/{space}/knowledge-bases/{kb_a}/documents/{broken['id']}/reparse",
        headers=auth,
    )
    check(
        "条2 reparse 受理(重置回队列)",
        rp.status_code == 200 and rp.json()["status"] == "pending",
        f"HTTP {rp.status_code} {rp.json().get('status')}",
    )

    # ---- 条 3:提问带引用;删除知识库再问 → 兜底话术不报错 ----
    ask = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/ask",
        json={
            "question": "混合检索用什么方式融合两路排名?参数 k 是多少?",
            "kb_ids": [kb_a],
            "top_k": 5,
        },
        headers=auth,
    )
    ev = sse_events(ask)
    types = [e["type"] for e in ev]
    cites = next((e["citations"] for e in ev if e["type"] == "citations"), [])
    answer = "".join(e.get("text", "") for e in ev if e["type"] == "delta")
    check(
        "条3 SSE 时序(meta…done,真模型)",
        types[0] == "meta" and types[-1] == "done" and "delta" in types,
        f"{types[:6]}… 共{len(types)}",
    )
    check(
        "条3 引用非空且可回链", bool(cites) and bool(cites[0].get("chunk_id")), f"{len(cites)} 条"
    )
    check("条3 答案有实质内容", len(answer) > 20, f"{len(answer)} 字:{answer[:40]}…")
    dl = ok_req(
        client, "DELETE", f"/api/v1/spaces/{space}/knowledge-bases/{kb_b}", headers=auth
    )
    check("条3 先删空知识库", dl.status_code == 204)
    ask2 = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/ask",
        json={"question": "删除之后还存在的内容有什么?", "kb_ids": [kb_a], "top_k": 5},
        headers=auth,
    )
    ev2 = sse_events(ask2)
    ans2 = "".join(e.get("text", "") for e in ev2 if e["type"] == "delta")
    ok2 = ev2[0]["type"] == "meta" and ev2[-1]["type"] == "done"
    check(
        "条3 删除后问答不报错且走兜底/正常分支",
        ok2,
        f"{[e['type'] for e in ev2]} 答案{len(ans2)}字",
    )
    # 删掉带内容的 A 库,再问 → 必须兜底话术
    ok_req(client, "DELETE", f"/api/v1/spaces/{space}/knowledge-bases/{kb_a}", headers=auth)
    ask3 = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/ask",
        json={"question": "现在还能引用到什么?", "top_k": 5},
        headers=auth,
    )
    ev3 = sse_events(ask3)
    ans3 = "".join(e.get("text", "") for e in ev3 if e["type"] == "delta")
    check(
        "条3 无资料时兜底话术(非报错)",
        ev3[-1]["type"] == "done" and "没有检索到" in ans3,
        ans3[:50],
    )

    # ---- 条 4:流中途断开 → 按 seq 续播不重不丢(重新上传一批文档) ----
    kb_c = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/knowledge-bases",
        json={"name": f"KB-C{tag}"},
        headers=auth,
    ).json()["id"]
    for fname, data in make_files().values():
        up_r = ok_req(
            client,
            "POST",
            f"/api/v1/spaces/{space}/knowledge-bases/{kb_c}/documents",
            files={"file": (fname, data)},
            headers=auth,
        )
        poll_doc(client, auth, space, kb_c, up_r.json()["id"])
    seen: list[dict] = []
    with client.stream(
        "POST",
        f"/api/v1/spaces/{space}/ask",
        json={
            "question": "请详细综述混合检索、RRF、重排、向量索引与全文检索的优缺点和适用场景。",
            "kb_ids": [kb_c],
            "top_k": 8,
        },
        headers=auth,
    ) as stream:
        for line in stream.iter_lines():
            if line.startswith("data: "):
                e = json.loads(line[6:])
                seen.append(e)
                if len([x for x in seen if x["type"] == "delta"]) >= 5:
                    break  # 模拟断网:直接掐断连接
    conv4 = seen[0]["conversation_id"]
    last_seq = seen[-1]["seq"]
    time.sleep(1)
    for _ in range(60):  # 等后台泵生成完(断连不终止生成,OPT-3)
        mr = client.get(
            f"/api/v1/spaces/{space}/conversations/{conv4}/messages", headers=auth
        )
        msgs = mr.json() if mr.status_code == 200 else []
        if len(msgs) >= 2:
            break
        time.sleep(2)
    resume = ok_req(
        client,
        "GET",
        f"/api/v1/spaces/{space}/conversations/{conv4}/stream",
        params={"after": last_seq},
        headers=auth,
    )
    rev = sse_events(resume)
    all_seq = [e["seq"] for e in seen] + [e["seq"] for e in rev]
    dedup = sorted(set(all_seq))
    check(
        "条4 断线续播可达",
        resume.status_code == 200
        and rev[0]["type"] in ("citations", "delta", "done", "meta", "tool_call"),
        f"HTTP {resume.status_code} 首帧 {rev[0]['type'] if rev else '∅'}",
    )
    check(
        "条4 seq 连续无重复(不重不丢)",
        dedup == list(range(dedup[0], dedup[0] + len(dedup))) and len(dedup) == len(all_seq),
        f"seq {dedup[0]}…{dedup[-1]} 共{len(dedup)}",
    )
    full = "".join(e.get("text", "") for e in seen + rev if e["type"] == "delta")
    persisted = msgs[-1]["content"] if len(msgs) >= 2 else ""
    check(
        "条4 续播全文与落库一致", full == persisted, f"流 {len(full)} 字 vs 库 {len(persisted)} 字"
    )

    # ---- 条 5:KB 范围受限 API Key(OPT-26:绑定成员 = 本人,门槛 Admin+) ----
    key_r = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/api-keys",
        json={
            "name": f"仅C库{tag}",
            "capabilities": ["chat"],
            "kb_ids": [kb_c],
            "member_user_id": me_id,
        },
        headers=auth,
    )
    plain = key_r.json()["plaintext"]
    kaf = {"X-API-Key": plain}
    s_in = ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/retrieval/search",
        json={"query": "RRF 融合", "kb_ids": [kb_c]},
        headers=kaf,
    )
    check(
        "条5 范围内检索可用",
        s_in.status_code == 200 and bool(s_in.json()),
        f"HTTP {s_in.status_code} 命中{len(s_in.json()) if s_in.status_code == 200 else 0}",
    )
    s_out = client.post(
        f"/api/v1/spaces/{space}/retrieval/search",
        json={"query": "RRF 融合", "kb_ids": [str(uuid.uuid4())]},
        headers=kaf,
    )
    check(
        "条5 范围外 403 API_KEY_SCOPE_DENIED",
        s_out.status_code == 403 and "API_KEY_SCOPE_DENIED" in s_out.text,
        f"HTTP {s_out.status_code}",
    )
    # 能力外路由:kb 列表属 DOCUMENTS 能力,chat-only Key 访问 → 403(授权表 fail-closed)。
    # 注:GET /api/v1/spaces 这类未接入 Key 认证的路由返回 401(与"没登录"同构),同样是拒绝。
    s_route = client.get(f"/api/v1/spaces/{space}/knowledge-bases", headers=kaf)
    check(
        "条5 能力外路由一律拒绝(fail-closed 403)",
        s_route.status_code == 403 and "API_KEY_CAPABILITY_DENIED" in s_route.text,
        f"HTTP {s_route.status_code}",
    )
    s_jwt_only = client.get("/api/v1/spaces", headers=kaf)
    check(
        "条5 未接 Key 认证的路由对 Key 一律 401",
        s_jwt_only.status_code == 401,
        f"HTTP {s_jwt_only.status_code}",
    )

    # ---- 条 6:Viewer 管理接口 403;非成员 404;审计 access.denied ----
    viewer = ok_req(
        client,
        "POST",
        "/api/v1/auth/register",
        json={"username": f"viewer{tag}", "nickname": "观众", "password": "secret-pass-1"},
    )
    v_auth = {"Authorization": f"Bearer {viewer.json()['access_token']}"}
    ok_req(
        client,
        "POST",
        f"/api/v1/spaces/{space}/members",
        json={"username": f"viewer{tag}", "role": 10},
        headers=auth,
    )
    mk = client.post(
        f"/api/v1/spaces/{space}/knowledge-bases", json={"name": "越权库"}, headers=v_auth
    )
    check(
        "条6 Viewer 建库 403",
        mk.status_code == 403 and "FORBIDDEN" in mk.text,
        f"HTTP {mk.status_code}",
    )
    kv = client.post(
        f"/api/v1/spaces/{space}/api-keys",
        json={
            "name": "越权key",
            "capabilities": ["chat"],
            "member_user_id": viewer.json()["user"]["id"],
        },
        headers=v_auth,
    )
    check(
        "条6 Viewer 建 Key 403(OPT-26 Admin+ 门槛)",
        kv.status_code == 403 and "FORBIDDEN" in kv.text,
        f"HTTP {kv.status_code}",
    )
    outsider = ok_req(
        client,
        "POST",
        "/api/v1/auth/register",
        json={"username": f"out{tag}", "nickname": "路人", "password": "secret-pass-1"},
    )
    o_auth = {"Authorization": f"Bearer {outsider.json()['access_token']}"}
    o_list = client.get(f"/api/v1/spaces/{space}/knowledge-bases", headers=o_auth)
    check("条6 非成员 404(防枚举)", o_list.status_code == 404, f"HTTP {o_list.status_code}")
    audit = ok_req(
        client, "GET", f"/api/v1/spaces/{space}/audit-logs", headers=auth
    ).json()
    denied = [a for a in audit["items"] if "denied" in a.get("action", "")]
    check("条6 审计有 access.denied 记录", bool(denied), f"{len(denied)} 条")

    print("\n=== 结果 ===")
    ok = all(r[1] for r in results)
    for name, passed, detail in results:
        if not passed:
            print(f"  {FAIL} {name} — {detail}")
    print(f"{len([r for r in results if r[1]])}/{len(results)} 通过")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
