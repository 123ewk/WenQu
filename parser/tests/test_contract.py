"""跨进程契约测试(基准 03):起真实 gRPC 服务,钉住 proto 契约与 Token 认证。

本文件属于「fake 拒绝区」——协议接缝必须真实现测试,不许 mock。
"""

from __future__ import annotations

from concurrent import futures

import grpc
import pytest

from parser_service.pb import parser_pb2, parser_pb2_grpc
from parser_service.server import _TOKEN_METADATA_KEY, ParserService, _TokenInterceptor

_TEST_TOKEN = "test-token"


@pytest.fixture()
def grpc_addr() -> str:
    server = grpc.server(
        futures.ThreadPoolExecutor(max_workers=2),
        interceptors=[_TokenInterceptor(_TEST_TOKEN)],
    )
    parser_pb2_grpc.add_ParserServiceServicer_to_server(ParserService(), server)
    port = server.add_insecure_port("localhost:0")
    server.start()
    yield f"localhost:{port}"
    server.stop(grace=None)


def _stub(addr: str) -> parser_pb2_grpc.ParserServiceStub:
    channel = grpc.insecure_channel(addr)
    return parser_pb2_grpc.ParserServiceStub(channel)


def test_parse_plain_text_returns_blocks(grpc_addr: str) -> None:
    stub = _stub(grpc_addr)
    resp = stub.Parse(
        parser_pb2.ParseRequest(
            document_id="doc-1",
            format="txt",
            content="第一段\n多行内容。\n\n第二段".encode(),
        ),
        metadata=((_TOKEN_METADATA_KEY, _TEST_TOKEN),),
    )
    assert resp.document_id == "doc-1"
    assert [b.text for b in resp.blocks] == ["第一段\n多行内容。", "第二段"]
    assert all(b.type == "paragraph" for b in resp.blocks)


def test_unsupported_format_is_invalid_argument(grpc_addr: str) -> None:
    stub = _stub(grpc_addr)
    with pytest.raises(grpc.RpcError) as exc_info:
        stub.Parse(
            parser_pb2.ParseRequest(document_id="doc-2", format="exe", content=b"MZ"),
            metadata=((_TOKEN_METADATA_KEY, _TEST_TOKEN),),
        )
    assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_parse_docx_end_to_end(grpc_addr: str) -> None:
    """协议接缝用真实 docx 字节走完整 gRPC 链路(fake 拒绝区)。"""
    import io

    from docx import Document

    buffer = io.BytesIO()
    doc = Document()
    doc.add_heading("合同模板", level=1)
    doc.add_paragraph("甲方与乙方约定如下条款。")
    doc.save(buffer)

    stub = _stub(grpc_addr)
    resp = stub.Parse(
        parser_pb2.ParseRequest(
            document_id="doc-3", format="docx", content=buffer.getvalue()
        ),
        metadata=((_TOKEN_METADATA_KEY, _TEST_TOKEN),),
    )
    assert resp.document_id == "doc-3"
    assert resp.blocks[0].type == "title" and resp.blocks[0].level == 1
    assert resp.blocks[0].text == "合同模板"
    assert any(b.type == "paragraph" and "甲方与乙方" in b.text for b in resp.blocks)


def test_dirty_document_is_data_loss(grpc_addr: str) -> None:
    stub = _stub(grpc_addr)
    with pytest.raises(grpc.RpcError) as exc_info:
        stub.Parse(
            parser_pb2.ParseRequest(
                document_id="doc-4", format="pdf", content=b"not-a-real-pdf"
            ),
            metadata=((_TOKEN_METADATA_KEY, _TEST_TOKEN),),
        )
    assert exc_info.value.code() == grpc.StatusCode.DATA_LOSS


def test_missing_token_is_unauthenticated(grpc_addr: str) -> None:
    stub = _stub(grpc_addr)
    with pytest.raises(grpc.RpcError) as exc_info:
        stub.Parse(parser_pb2.ParseRequest(document_id="doc-3", format="txt", content=b"x"))
    assert exc_info.value.code() == grpc.StatusCode.UNAUTHENTICATED


def test_wrong_token_is_unauthenticated(grpc_addr: str) -> None:
    stub = _stub(grpc_addr)
    with pytest.raises(grpc.RpcError) as exc_info:
        stub.Parse(
            parser_pb2.ParseRequest(document_id="doc-4", format="txt", content=b"x"),
            metadata=((_TOKEN_METADATA_KEY, "wrong-token"),),
        )
    assert exc_info.value.code() == grpc.StatusCode.UNAUTHENTICATED
